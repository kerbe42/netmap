"""Check a scan against the asset list you were handed.

Whoever takes over a network usually gets a spreadsheet of what is supposed to be on it.
This matches its rows to what was actually found - by management address, then serial
number, then name, then MAC - and sorts everything into: listed and found, listed but not
found, found but not listed, and found but disagreeing with the list (a different serial,
model or name at the same address).
"""
from __future__ import annotations

import csv
import ipaddress
import re
from dataclasses import dataclass, field
from typing import Optional

from .model import Inventory
from .util import norm_mac, short_name

FIELDS = {
    "ip": ["ip", "ip address", "ipaddress", "ipv4", "management ip", "mgmt ip", "mgmt", "address", "ip addr"],
    "name": ["name", "hostname", "host name", "host", "device", "device name", "asset name", "sysname", "system name"],
    "serial": ["serial", "serial number", "serialnumber", "serial no", "s/n", "sn"],
    "mac": ["mac", "mac address", "macaddress", "hw address", "hardware address"],
    "model": ["model", "model number", "product", "part number", "pid"],
    "site": ["site", "location", "building", "room"],
}


def read_table(path: str) -> tuple[list[str], list[list[str]]]:
    """Headers and rows from a CSV (any common delimiter) or the first sheet of an .xlsx."""
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [["" if c is None else str(c).strip() for c in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        with open(path, encoding="utf-8-sig", errors="replace", newline="") as f:
            text = f.read()
        rows = [[c.strip() for c in r] for r in csv.reader(text.splitlines(), delimiter=_delimiter(text))]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return [], []
    # the header is the first row that names at least one field we know; anything above it is a title
    for i, r in enumerate(rows[:10]):
        if guess_columns(r):
            return r, rows[i + 1 :]
    return rows[0], rows[1:]


def _delimiter(text: str) -> str:
    """The separator most lines agree on. csv.Sniffer gives up when a title line sits above
    the header, which is how handover spreadsheets usually arrive."""
    lines = [ln for ln in text.splitlines()[:30] if ln.strip()]
    best, best_score = ",", -1
    for d in (",", ";", "\t", "|"):
        counts = [ln.count(d) for ln in lines]
        used = [c for c in counts if c]
        if not used:
            continue
        mode = max(set(used), key=used.count)
        score = used.count(mode) * mode
        if score > best_score:
            best, best_score = d, score
    return best


def guess_columns(headers: list[str]) -> dict[str, int]:
    out: dict[str, int] = {}
    norm = [re.sub(r"[\s_\-.]+", " ", (h or "").strip().lower()) for h in headers]
    for fname, names in FIELDS.items():
        for i, h in enumerate(norm):
            if h in names and i not in out.values():
                out[fname] = i
                break
    return out


@dataclass
class Match:
    row: int  # index into the list's rows
    listed: dict  # the row's recognised fields
    node: str = ""  # what it matched, "" if nothing
    how: str = ""  # address | serial | name | mac
    differences: list[str] = field(default_factory=list)


@dataclass
class Reconciliation:
    matches: list[Match]
    unlisted: list[str]  # devices found that no row matched

    @property
    def found(self):
        return [m for m in self.matches if m.node and not m.differences]

    @property
    def missing(self):
        return [m for m in self.matches if not m.node]

    @property
    def differ(self):
        return [m for m in self.matches if m.node and m.differences]


def _ip(v: str) -> Optional[str]:
    v = (v or "").strip().split("/")[0]
    try:
        return str(ipaddress.ip_address(v))
    except ValueError:
        return None


def reconcile(inv: Inventory, rows: list[list[str]], columns: dict[str, int]) -> Reconciliation:
    by_serial: dict[str, str] = {}
    by_name: dict[str, str] = {}
    by_mac: dict[str, str] = {}
    for d in inv.devices.values():
        if d.serial:
            by_serial.setdefault(_norm_serial(d.serial), d.id)
        for c in d.components:
            if c.serial:
                by_serial.setdefault(_norm_serial(c.serial), d.id)
        for n in (d.name, d.dns_name, inv.note(d.id).get("name", "")):
            if n:
                by_name.setdefault(short_name(n), d.id)
        for m in d.macs + [d.lldp_chassis_id]:
            if m:
                by_mac.setdefault(m.lower(), d.id)
    for ip, h in inv.hosts.items():
        if ip in inv.ip_to_device:
            continue
        for n in (h.hostname, inv.note(ip).get("name", "")):
            if n:
                by_name.setdefault(short_name(n), ip)
        if h.mac:
            by_mac.setdefault(h.mac.lower(), ip)

    def get(r, f):
        i = columns.get(f)
        return r[i].strip() if i is not None and i < len(r) and r[i] else ""

    matches: list[Match] = []
    hit_devices: set[str] = set()
    for n, r in enumerate(rows):
        listed = {f: get(r, f) for f in FIELDS if get(r, f)}
        if not listed:
            continue
        m = Match(row=n, listed=listed)
        ip = _ip(listed.get("ip", ""))
        if ip and (ip in inv.ip_to_device or ip in inv.hosts):
            m.node, m.how = inv.ip_to_device.get(ip, ip), "address"
        elif _norm_serial(listed.get("serial", "")) in by_serial:
            m.node, m.how = by_serial[_norm_serial(listed["serial"])], "serial"
        elif listed.get("name") and short_name(listed["name"]) in by_name:
            m.node, m.how = by_name[short_name(listed["name"])], "name"
        elif norm_mac(listed.get("mac")) and norm_mac(listed["mac"]) in by_mac:
            m.node, m.how = by_mac[norm_mac(listed["mac"])], "mac"
        if m.node:
            d = inv.devices.get(m.node)
            if d is not None:
                hit_devices.add(d.id)
                # matched the same way it was indexed: stripped and upper-cased on both sides
                serials = {_norm_serial(d.serial)} | {_norm_serial(c.serial) for c in d.components if c.serial}
                if listed.get("serial") and _norm_serial(listed["serial"]) not in serials and d.serial:
                    m.differences.append(f"serial: list {listed['serial']}, device {d.serial}")
                if listed.get("model") and d.model and not _same_model(listed["model"], d.model, d.components):
                    m.differences.append(f"model: list {listed['model']}, device {d.model}")
                known_names = {short_name(d.name), short_name(d.dns_name or ""), short_name(inv.note(d.id).get("name", ""))} - {""}
                if listed.get("name") and known_names and short_name(listed["name"]) not in known_names:
                    m.differences.append(f"name: list {listed['name']}, device {inv.display_name(d.id)}")
                if ip and m.how != "address":
                    m.differences.append(f"address: list {ip}, found at {d.id}")
            else:
                h = inv.hosts.get(m.node)
                if h is not None and listed.get("mac") and h.mac and norm_mac(listed["mac"]) != h.mac:
                    m.differences.append(f"MAC: list {listed['mac']}, seen {h.mac}")
                if h is not None and _listed_as_kit(listed, h):
                    # the list says this is a managed device (it has a serial/model, or it
                    # announced itself as network kit) but only its address answered
                    m.differences.append(f"no SNMP answer: {m.node} was seen on the network but could not be polled")
        matches.append(m)
    unlisted = sorted(set(inv.devices) - hit_devices, key=lambda x: ipaddress.ip_address(x))
    return Reconciliation(matches, unlisted)


def _norm_serial(s: str) -> str:
    return (s or "").strip().upper()


_KIT_ROLES = {"switch", "l3switch", "router", "firewall", "wireless", "unpolled"}


def _listed_as_kit(listed: dict, host) -> bool:
    """Does the list describe network kit we should have polled? A row carrying a serial
    or model, or a host that announced itself as a switch/router/AP, or one an SNMP probe
    was tried on - as opposed to a printer or PC listed for completeness."""
    return bool(listed.get("serial") or listed.get("model")) or getattr(host, "snmp_failed", False) or getattr(host, "role", "") in _KIT_ROLES


def _same_model(listed: str, model: str, components) -> bool:
    a = re.sub(r"[^a-z0-9]", "", listed.lower())
    cands = [model] + [c.model for c in components if c.model]
    return any(a and (a in re.sub(r"[^a-z0-9]", "", c.lower()) or re.sub(r"[^a-z0-9]", "", c.lower()) in a) for c in cands if c)


def write_csv(inv: Inventory, rec: Reconciliation, path: str) -> str:
    from .graph import SafeCsvWriter

    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = SafeCsvWriter(f)  # names/serials off the list or the network are text, never formulas
        w.writerow(["result", "list name", "list address", "list serial", "matched", "matched by", "found name", "differences"])
        for m in rec.matches:
            state = "not found" if not m.node else ("differs" if m.differences else "found")
            w.writerow([state, m.listed.get("name", ""), m.listed.get("ip", ""), m.listed.get("serial", ""), m.node, m.how,
                        inv.display_name(m.node) if m.node else "", "; ".join(m.differences)])
        for did in rec.unlisted:
            d = inv.devices[did]
            w.writerow(["not in list", "", "", "", did, "", d.name, f"{d.vendor} {d.model} {d.serial}".strip()])
    return path
