"""Excel workbook of the inventory: the hand-over artefact for documenting a network.

One sheet per thing you get asked about - what is on the network and what hardware it
is built from (serials, supplies, optics for the asset register), how it is wired,
which addresses are in use, what VLANs exist, and what we saw but could not get into.
Everything here comes from the saved map, so it can be rebuilt without touching the
network again.
"""
from __future__ import annotations

import ipaddress
import logging
import time
from collections import Counter

from .graph import edge_ports, ipam_rows, vlan_rows
from .model import Inventory

log = logging.getLogger("netmap.report")

HEADER_FILL = "FF1F3864"
BANDS = {"device": "FFDDEBF7", "host": "FFF2F2F2", "subnet": "FFE2EFDA"}


def _sheet(wb, title, headers, rows, widths=None, freeze="A2"):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=HEADER_FILL)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = freeze
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for i, h in enumerate(headers, 1):
        w = (widths or {}).get(h)
        if w is None:
            longest = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:400] if len(r) >= i] or [8])
            w = min(max(longest + 2, 9), 52)
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def export_xlsx(inv: Inventory, g, path: str) -> str:
    """Write the whole inventory as a multi-sheet workbook. Returns the path written."""
    try:
        from openpyxl import Workbook
    except ImportError as e:  # pragma: no cover - dependency is declared, but say so clearly
        raise RuntimeError("openpyxl is required for --xlsx (pip install openpyxl)") from e

    wb = Workbook()
    wb.remove(wb.active)

    roles = Counter(d.role for d in inv.devices.values())
    vendors = Counter(d.vendor or "unknown" for d in inv.devices.values())
    host_roles = Counter(h.role for ip, h in inv.hosts.items() if ip not in inv.ip_to_device)
    edge_kinds = Counter(a["kind"] for _, _, a in g.edges(data=True))
    unpolled = [(n, a) for n, a in g.nodes(data=True) if a.get("role") == "unpolled"]
    ipam = ipam_rows(inv)

    summary = [
        ["Generated", time.strftime("%Y-%m-%d %H:%M:%S")],
        ["SNMP devices", len(inv.devices)],
        ["Hosts (no SNMP)", sum(1 for ip in inv.hosts if ip not in inv.ip_to_device)],
        ["Subnets", len(inv.subnets)],
        ["VLANs", len(vlan_rows(inv))],
        ["Addresses probed without an SNMP answer", len(inv.unreachable)],
        ["Neighbours seen but not polled", len(unpolled)],
        ["Addresses in use / usable", f"{sum(r['used'] for r in ipam)} / {sum(r['usable'] for r in ipam)}"],
        [],
        ["Devices by role"] + [],
    ]
    summary += [[f"  {r}", n] for r, n in roles.most_common()]
    summary += [[], ["Devices by vendor"]] + [[f"  {v}", n] for v, n in vendors.most_common()]
    summary += [[], ["Hosts by type"]] + [[f"  {r}", n] for r, n in host_roles.most_common()]
    summary += [[], ["Links by kind"]] + [[f"  {k}", n] for k, n in edge_kinds.most_common()]
    _sheet(wb, "Summary", ["Item", "Value"], summary, widths={"Item": 44, "Value": 30}, freeze="A2")

    _sheet(
        wb, "Devices",
        ["IP", "Name", "Role", "Vendor", "Model", "OS version", "Serial", "Location", "Contact", "OS / sysDescr", "All IPs",
         "Interfaces", "VLANs", "LLDP", "CDP", "ARP", "Routes", "FDB", "Uptime (days)", "Depth", "Discovered via", "Credential", "Errors"],
        [
            [d.id, d.name, d.role, d.vendor, d.model, d.os_version, d.serial, d.location, d.contact, d.sysdescr[:300], " ".join(d.ips),
             len(d.interfaces), len(d.vlans), sum(n.proto == "lldp" for n in d.neighbors), sum(n.proto == "cdp" for n in d.neighbors),
             len(d.arp), len(d.routes), len(d.fdb), d.uptime_s // 86400, d.depth, d.discovered_via, d.credential, "; ".join(d.errors)[:300]]
            for d in sorted(inv.devices.values(), key=lambda x: (x.depth, ipaddress.ip_address(x.id)))
        ],
        widths={"OS / sysDescr": 50, "Errors": 40},
    )

    _sheet(
        wb, "IPAM",
        ["Subnet", "Size", "Usable", "In use", "Free", "Utilisation %", "VLAN", "Gateways", "Discovered from", "Swept"],
        [[r["cidr"], r["size"], r["usable"], r["used"], r["free"], r["utilisation_pct"], r["vlan"], r["gateways"], r["sources"], "yes" if r["swept"] else "no"] for r in ipam],
    )

    _sheet(
        wb, "VLANs", ["VLAN", "Name(s)", "Devices", "Carried on"],
        [[vid, " / ".join(sorted(names)), len(devs), " ".join(sorted(devs))] for vid, (names, devs) in sorted(vlan_rows(inv).items())],
        widths={"Carried on": 50},
    )

    _sheet(
        wb, "Links",
        ["A", "A name", "A port", "B", "B name", "B port", "Kind", "Detail"],
        [
            [u, g.nodes[u].get("label"), pu, v, g.nodes[v].get("label"), pv, a["kind"], a.get("label", "")]
            for u, v, a in g.edges(data=True) if a.get("kind") in ("lldp", "cdp", "l3")
            for pu, pv in [edge_ports(u, v, a)]
        ],
    )

    _sheet(
        wb, "Hosts",
        ["IP", "Hostname", "MAC", "Vendor (OUI)", "Type", "Subnet", "Switch", "Port", "VLAN", "Open ports", "Seen from", "SNMP tried"],
        [
            [
                ip, h.hostname, h.mac or "", h.vendor, h.role, inv.subnet_for_ip(ip) or "",
                (inv.devices[f["device"]].name or f["device"]) if (f := next((s for s in h.seen_on if s["via"] == "fdb"), None)) and f["device"] in inv.devices else (f["device"] if f else ""),
                f["interface"] if f else "", f["vlan"] if f else "",
                " ".join(f"{x['port']}/{x['service']}" for x in h.ports), " ".join(h.sources), "no answer" if h.snmp_failed else "",
            ]
            for ip, h in sorted(inv.hosts.items(), key=lambda kv: ipaddress.ip_address(kv[0])) if ip not in inv.ip_to_device
        ],
        widths={"Open ports": 40},
    )

    _sheet(
        wb, "Interfaces",
        ["Device", "Device name", "ifIndex", "Name", "Description", "Alias", "MAC", "Speed (Mbps)", "Admin", "Oper", "VLAN", "Mode", "LAG", "Addresses"],
        [
            [d.id, d.name, i.index, i.name, i.descr, i.alias, i.mac or "", i.speed_mbps, "up" if i.admin_up else "down", "up" if i.oper_up else "down",
             i.vlan if i.vlan is not None else "", i.mode, i.lag, " ".join(i.ips)]
            for d in inv.devices.values() for i in d.interfaces
        ],
        widths={"Description": 34, "Alias": 30},
    )

    _sheet(
        wb, "Hardware",
        ["Device", "Device name", "Class", "Name", "Description", "Model", "Serial", "HW rev", "FW rev", "SW rev", "FRU"],
        [
            [d.id, d.name, c.cls, c.name, c.descr, c.model, c.serial, c.hw_rev, c.fw_rev, c.sw_rev, "yes" if c.fru else "no"]
            for d in inv.devices.values() for c in d.components
        ],
        widths={"Description": 40},
    )

    _sheet(
        wb, "Gaps",
        ["What", "Identifier", "Detail", "Why it matters"],
        [["Unpolled neighbour", a.get("label"), f"ip={a.get('ip') or '-'} chassis={a.get('chassis_id', '') or '-'} {a.get('model', '')[:60]}",
          "Announced by a neighbour but no credentials worked - an unmanaged device or one outside the handover"] for _, a in unpolled]
        + [["No SNMP answer", ip, f"first seen via {via}", "Probed inside scope and never answered - host, filtered, or different credentials"]
           for ip, via in sorted(inv.unreachable.items(), key=lambda kv: ipaddress.ip_address(kv[0]))[:2000]]
        + [["Subnet never swept", r["cidr"], f"gateways: {r['gateways'] or '-'}",
            "Addresses here are known only from ARP/routes, so utilisation is a floor not a count"] for r in ipam if not r["swept"]],
        widths={"Detail": 46, "Why it matters": 60},
    )

    wb.save(path)
    log.info("wrote %s (%d sheets)", path, len(wb.sheetnames))
    return path
