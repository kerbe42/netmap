import asyncio
import ipaddress
import json
import os
import re
import tempfile

import pytest

from netmap import oids as O
from netmap.collect import MAX_COMPONENTS, CollectOptions, collect_device, collect_entity, collect_system
from netmap.crawl import CrawlConfig, Crawler
from netmap.graph import build_graph, enrich_inventory, export_csv, export_dot, export_graphml, ipam_rows, text_summary, vlan_rows
from netmap.model import Device, Interface, Inventory
from netmap.render import render_html
from netmap.report import export_xlsx
from netmap.snmp import Credential
from netmap.util import in_scope, mask_to_prefix, oid_suffix, oui_vendor, parse_os_version, portlist_ports, short_name

from . import labnet
from .fake_snmp import FakeSession, make_prober


def crawl(base="10", **kw):
    lab = labnet.build(base)
    tables = {ip: d.values() for ip, d in lab.items()}
    prober = make_prober(tables)
    cfg = CrawlConfig(seeds=[f"{base}.0.0.1"], credentials=[Credential(kind="v2c", community="lab", label="lab")], scope=[ipaddress.ip_network(f"{base}.0.0.0/8")], workers=4, **kw)
    inv = Inventory()
    c = Crawler(cfg, inv, engine=object(), prober=prober)
    asyncio.run(c.run())
    return inv, prober, lab


def test_util():
    assert oid_suffix("1.2.3.4.5", "1.2.3") == [4, 5]
    assert oid_suffix("1.2.30.4", "1.2.3") == []
    assert mask_to_prefix("255.255.255.0") == 24
    assert short_name("dist-sw1.example.test") == "dist-sw1"
    assert short_name("core-rtr.example.test(FDO123)") == "core-rtr"


def test_crawl_finds_all_devices_and_links():
    inv, prober, lab = crawl()
    assert set(inv.devices) == {"10.0.0.1", "10.0.0.2", "10.1.0.2"}
    r1, sw1, sw2 = inv.devices["10.0.0.1"], inv.devices["10.0.0.2"], inv.devices["10.1.0.2"]
    assert (r1.name, r1.vendor, r1.role, r1.model, r1.serial) == ("core-rtr", "Cisco", "router", "ISR4331/K9", "FDO2222R1XX")
    assert (sw1.role, sw2.role, sw2.vendor) == ("l3switch", "switch", "HP")
    assert sw1.depth == 1 and sw2.depth == 2
    # SW1 reached by CDP mgmt address from R1; SW2 by LLDP mgmt address from SW1
    assert sw1.discovered_via.startswith("cdp:") and sw2.discovered_via.startswith("lldp:")
    # the same box reached via its SVI address is deduped, not re-collected
    assert inv.ip_to_device["10.1.0.1"] == "10.0.0.2"
    assert "10.1.0.1" not in inv.devices
    # out-of-scope neighbour and WAN next-hop were never probed
    assert "192.168.99.2" not in prober.probed and "192.168.99.7" not in inv.hosts and "203.0.113.1" not in prober.probed
    # interface details
    gi1 = sw1.iface(1)
    assert gi1.name == "Gi1/0/1" and gi1.alias == "uplink core-rtr" and gi1.speed_mbps == 1000 and "10.0.0.2/30" in gi1.ips
    assert sw1.vlans == {10: "USERS", 20: "SERVERS", 30: "BRANCH"}
    # neighbours parsed
    lldp = {n.remote_name: n for n in sw1.neighbors if n.proto == "lldp"}
    assert lldp["acc-sw2"].remote_mgmt_ips == ["10.1.0.2"] and lldp["acc-sw2"].local_port == "Gi1/0/2" and lldp["acc-sw2"].local_if_index == 2
    assert lldp["ap-lobby"].remote_chassis_id == labnet.MAC_AP and lldp["ap-lobby"].remote_caps == "wlan-ap"
    cdp = [n for n in sw1.neighbors if n.proto == "cdp"][0]
    assert cdp.remote_mgmt_ips == ["10.0.0.1"] and cdp.local_port == "Gi1/0/1" and cdp.remote_port == "GigabitEthernet0/0/0"
    # routes / arp / fdb
    assert any(r.dest == "0.0.0.0/0" and r.nexthop == "10.0.0.1" for r in sw1.routes)
    assert any(r.dest == "10.1.0.0/24" and r.nexthop == "10.0.0.2" for r in r1.routes)
    assert {a.ip for a in sw1.arp} >= {"10.1.0.50", "10.1.0.51", "10.2.0.10", "10.1.0.60"}
    assert any(f.mac == labnet.MAC_C and f.if_index == 5 and f.vlan == 20 for f in sw1.fdb)
    assert any(f.mac == labnet.MAC_A and f.if_index == 3 and f.vlan is None for f in sw2.fdb)
    # hosts from ARP; the phone was LLDP-announced with a mgmt IP that does not speak SNMP
    assert "10.1.0.50" in inv.hosts and inv.hosts["10.1.0.50"].mac == labnet.MAC_A
    assert inv.hosts["10.1.0.90"].snmp_failed is True
    # subnets: three in scope plus the out-of-scope one recorded from the SVI (not crawled)
    assert {"10.0.0.0/30", "10.1.0.0/24", "10.2.0.0/24"} <= set(inv.subnets)
    assert not sw1.errors and not sw2.errors and not r1.errors


def test_hardware_os_version_port_vlans_and_lags():
    inv, _, _ = crawl()
    r1, sw1, sw2 = inv.devices["10.0.0.1"], inv.devices["10.0.0.2"], inv.devices["10.1.0.2"]
    assert (r1.os_version, sw1.os_version, sw2.os_version) == ("17.6.4", "16.12.4", "YA.16.10.0016")
    # ENTITY-MIB: the stack, both members, their supplies, a fan, the uplink module and its optic.
    # Slot containers, the sensor, the copper port, the empty SFP cage and the anonymous entity are gone.
    comps = {c.index: c for c in sw1.components}
    assert sorted(comps) == [1, 1000, 1002, 1003, 1005, 1008, 2000, 2002]
    assert [(c.name, c.serial) for c in sw1.components if c.cls == "chassis"] == [("Switch 1", "FOC1234SW1X"), ("Switch 2", "FOC1234SW2Y")]
    sfp = comps[1008]
    assert (sfp.cls, sfp.model, sfp.serial, sfp.hw_rev, sfp.fru) == ("port", "SFP-10G-SR", "AVD2045K1LM", "V03", True)
    # parents skip the dropped cage/slot: the optic hangs off the module, the supply off its chassis
    assert (sfp.parent, comps[1002].parent, comps[1005].parent, comps[1000].parent, comps[1].parent) == (1005, 1000, 1000, 1, 0)
    assert (comps[1].cls, comps[1].fru, comps[1003].cls, comps[1000].sw_rev, comps[1000].fw_rev) == ("stack", False, "fan", "16.12.4", "16.12.2r")
    # the device itself still carries the primary chassis
    assert (sw1.model, sw1.serial) == ("WS-C3850-24T", "FOC1234SW1X")
    assert [(c.cls, c.model, c.serial) for c in r1.components] == [("chassis", "ISR4331/K9", "FDO2222R1XX")]
    # Cisco: trunk with its native VLAN, access ports with theirs, the routed port untouched
    assert [(sw1.iface(i).mode, sw1.iface(i).vlan) for i in (2, 5, 24, 1, 10)] == [("trunk", 1), ("access", 20), ("access", 10), ("", None), ("", None)]
    # Q-BRIDGE on the HP: PVID per bridge port; tagged in two VLANs = trunk, one tagged voice VLAN is not
    assert {p: (sw2.iface(p).vlan, sw2.iface(p).mode) for p in (1, 2, 3, 4, 24, 289, 21)} == {
        1: (1, "access"), 2: (1, "access"), 3: (10, "access"), 4: (10, "access"), 24: (1, "trunk"), 289: (20, "access"), 21: (None, "")}
    # LAG: members name their aggregator; 0 and self-references mean "not aggregated"
    assert [sw2.iface(p).lag for p in (21, 22, 24, 3, 289)] == ["Trk1", "Trk1", "", "", ""]
    assert not any(i.lag for d in (r1, sw1) for i in d.interfaces)
    # ifLastChange is TimeTicks
    assert sw1.iface(5).last_change_s == 987 and sw2.iface(3).last_change_s == 13 and r1.iface(2).last_change_s == 12


def test_switch_tables_are_walked_only_where_they_can_exist():
    """Routers and servers must not pay for VLAN walks, and the bridge-port map is walked once."""
    lab = labnet.build("10")
    walked = {}
    for ip, d in lab.items():
        s = FakeSession(ip, d.values())
        dev = asyncio.run(collect_device(s, ip, CollectOptions()))
        assert not dev.errors, dev.errors
        walked[dev.name] = s.walked
    port_tables = {O.DOT1Q_PVID, O.DOT1Q_VLAN_CUR_EGRESS, O.CISCO_VM_VLAN, O.CISCO_TRUNK_STATUS, O.CISCO_TRUNK_NATIVE}
    assert not port_tables & set(walked["core-rtr"])
    assert O.CISCO_VM_VLAN in walked["dist-sw1"] and O.DOT1Q_PVID not in walked["dist-sw1"]  # Cisco answered its own MIB
    assert O.DOT1Q_PVID in walked["acc-sw2"] and O.CISCO_VM_VLAN not in walked["acc-sw2"]
    for name, w in walked.items():  # dot1q FDB (dist-sw1) and dot1d FDB (acc-sw2) both reuse the one walk
        assert w.count(O.DOT1D_BASE_PORT_IFINDEX) == 1, name
        assert O.LAG_ATTACHED_AGG in w and O.IF_LAST_CHANGE in w


def test_entity_component_cap_keeps_the_boxes_before_the_optics():
    t = {f"{O.ENT_CLASS}.1": 3, f"{O.ENT_MODEL}.1": "N7K-C7018", f"{O.ENT_SERIAL}.1": "JAF0000X01"}
    for n in range(2, 2 + MAX_COMPONENTS + 100):  # more optics than the cap, then a supply at the end
        t[f"{O.ENT_CLASS}.{n}"], t[f"{O.ENT_SERIAL}.{n}"], t[f"{O.ENT_CONTAINED_IN}.{n}"] = 10, f"SFP{n:05d}", 1
    last = 2 + MAX_COMPONENTS + 100
    t[f"{O.ENT_CLASS}.{last}"], t[f"{O.ENT_MODEL}.{last}"], t[f"{O.ENT_CONTAINED_IN}.{last}"] = 6, "N7K-AC-6.0KW", 1
    dev = Device(id="10.9.9.9")
    asyncio.run(collect_entity(FakeSession(dev.id, t), dev))
    assert len(dev.components) == MAX_COMPONENTS and not dev.errors
    assert [c.cls for c in dev.components][:1] == ["chassis"] and dev.components[-1].model == "N7K-AC-6.0KW"
    assert all(c.parent == 1 for c in dev.components[1:]) and (dev.model, dev.serial) == ("N7K-C7018", "JAF0000X01")


@pytest.mark.parametrize(
    "sysdescr, vendor, expected",
    [
        ("Cisco IOS Software, IOS-XE Software, Catalyst L3 Switch Software (CAT3K_CAA-UNIVERSALK9-M), Version 16.12.4, RELEASE SOFTWARE (fc5)", "Cisco", "16.12.4"),
        ("Cisco IOS Software [Cupertino], Catalyst L3 Switch Software (CAT9K_IOSXE), Version 17.9.4a, RELEASE SOFTWARE (fc5)", "Cisco", "17.9.4a"),
        ("Cisco IOS Software, C2960X Software (C2960X-UNIVERSALK9-M), Version 15.2(7)E4, RELEASE SOFTWARE (fc2)\r\nTechnical Support: http://www.cisco.com/techsupport", "Cisco", "15.2(7)E4"),
        ("Cisco Internetwork Operating System Software \r\nIOS (tm) C2950 Software (C2950-I6Q4L2-M), Version 12.1(22)EA14, RELEASE SOFTWARE (fc1)", "Cisco", "12.1(22)EA14"),
        ("Cisco NX-OS(tm) n9000, Software (n9000-dk9), Version 9.3(8), RELEASE SOFTWARE Copyright (c) 2002-2021 by Cisco Systems, Inc.", "Cisco", "9.3(8)"),
        ("Cisco NX-OS(tm) nxos.7.0.3.I7.9.bin, Software (nxos), Version 7.0(3)I7(9), RELEASE SOFTWARE", "Cisco", "7.0(3)I7(9)"),
        ("Cisco Adaptive Security Appliance Version 9.16(3)", "Cisco", "9.16(3)"),
        ("Cisco IOS XR Software (Cisco ASR9K Series),  Version 6.5.3[Default]\nCopyright (c) 2019 by Cisco Systems, Inc.", "Cisco", "6.5.3"),
        ("Juniper Networks, Inc. ex4300-48p Ethernet Switch, kernel JUNOS 20.4R3-S2, Build date: 2022-01-20 03:47:45 UTC", "Juniper", "20.4R3-S2"),
        ("Juniper Networks, Inc. mx480 internet router, kernel JUNOS 21.2R3.8, Build date: 2022-03-01", "Juniper", "21.2R3.8"),
        ("Arista Networks EOS version 4.28.3M running on an Arista Networks DCS-7050SX3-48YC8", "Arista", "4.28.3M"),
        ("HP J9772A 2530-48G-PoEP Switch, revision YA.16.10.0016, ROM YA.15.20 (/ws/swbuildm/rel_yakima_qaoff/code/build/lakes)", "HP", "YA.16.10.0016"),
        ("ProCurve J8697A Switch 5406zl, revision K.15.18.0013, ROM K.15.30 (/sw/code/build/btm(K_15))", "HP", "K.15.18.0013"),
        ("Aruba JL658A 6300M 24SFP+ 4SFP56 Swch FL.10.08.1010", "HP", "FL.10.08.1010"),
        ("ArubaOS-CX Version: FL.10.08.1010", "Aruba", "FL.10.08.1010"),
        ("ArubaOS (MODEL: 7010), Version 8.10.0.2 (84785)", "Aruba", "8.10.0.2"),
        ("FortiGate-60F v7.2.5,build1517,230606 (GA.F)", "Fortinet", "7.2.5"),
        ("FortiGate-60F", "Fortinet", ""),
        ("Palo Alto Networks PA-3220 series firewall", "Palo Alto", ""),
        ("Palo Alto Networks PAN-OS 10.2.4-h4", "Palo Alto", "10.2.4-h4"),
        ("RouterOS 7.12", "MikroTik", "7.12"),
        ("RouterOS RB4011iGS+", "MikroTik", ""),  # a model, not a version
        ("EdgeOS v2.0.9-hotfix.6.5574651.221230.1015", "Ubiquiti", "2.0.9-hotfix.6"),
        ("USW-24-PoE, 6.5.59.14777, Linux 3.6.5", "Ubiquiti", "6.5.59.14777"),
        ("EdgeSwitch 24-Port Lite, 1.9.3.5381055, Linux 3.6.5-1b604f2a", "Ubiquiti", "1.9.3.5381055"),
        ("SonicWALL TZ 370 (SonicOS 7.0.1-5035)", "SonicWall", "7.0.1-5035"),
        ("VMware ESXi 7.0.3 build-20036589 VMware, Inc. x86_64", "", "7.0.3 build-20036589"),
        ("pfSense fw.example.test 2.7.0-RELEASE FreeBSD 14.0-CURRENT amd64", "FreeBSD/pfSense", "2.7.0-RELEASE"),
        ("Dell EMC Networking OS10 Enterprise.\r\nSystem Description: OS10 Enterprise.\r\nOS Version: 10.5.2.6.\r\nSystem Type: S4148F-ON", "Dell/Force10", "10.5.2.6"),
        ("Linux host 5.15.0-91-generic #101-Ubuntu SMP Tue Nov 14 13:30:08 UTC 2023 x86_64", "Net-SNMP", "5.15.0-91-generic"),
        ("Linux host 5.15.0-91-generic #101-Ubuntu SMP Tue Nov 14 13:30:08 UTC 2023 x86_64", "", "5.15.0-91-generic"),
        ("Linux DiskStation 4.4.59+ #25426 SMP PREEMPT Mon Dec 14 18:48:50 CST 2020 x86_64", "Synology", ""),  # kernel is not DSM
        ("Hardware: Intel64 Family 6 Model 85 Stepping 7 AT/AT COMPATIBLE - Software: Windows Version 6.3 (Build 17763 Multiprocessor Free)", "Microsoft", "6.3 Build 17763"),
        ("Some Appliance Firmware Version 3.2.1 (b17)", "", "3.2.1"),
        ("HP ETHERNET MULTI-ENVIRONMENT,ROM none,JETDIRECT,JD153,EEPROM JSI24090012,CIDATE 07/02/2021", "HP", ""),
        ("Cisco Controller", "Cisco", ""),
        ("", "", ""),
    ],
)
def test_parse_os_version(sysdescr, vendor, expected):
    assert parse_os_version(sysdescr, vendor) == expected


def test_os_version_from_vendor_mib_and_offline():
    # FortiGate's sysDescr is just the model: one GET to fgSysVersion fills the gap
    t = {O.SYS_DESCR: b"FortiGate-60F", O.SYS_OBJECTID: "1.3.6.1.4.1.12356.101.1.60", O.SYS_NAME: b"fw1",
         O.OS_VERSION_OIDS["Fortinet"]: b"v7.2.5,build1517,230606 (GA.F)"}
    dev = Device(id="10.9.9.1")
    s = FakeSession(dev.id, t)
    asyncio.run(collect_system(s, dev))
    assert (dev.vendor, dev.os_version) == ("Fortinet", "7.2.5") and not dev.errors
    # ...and nobody else is asked for it
    s2, dev2 = FakeSession("10.9.9.2", {O.SYS_DESCR: b"Cisco IOS Software, Version 15.2(7)E4", O.SYS_OBJECTID: "1.3.6.1.4.1.9.1.1"}), Device(id="10.9.9.2")
    asyncio.run(collect_system(s2, dev2))
    assert dev2.os_version == "15.2(7)E4" and not set(O.OS_VERSION_OIDS.values()) & set(s2.got)
    # a map saved before os_version existed gains it when it is loaded and graphed
    inv = Inventory.from_dict({"devices": {"10.9.9.3": {"id": "10.9.9.3", "sysdescr": "Cisco Adaptive Security Appliance Version 9.16(3)23", "vendor": "Cisco"}}})
    enrich_inventory(inv)
    assert inv.devices["10.9.9.3"].os_version == "9.16(3)23"


def test_portlist_ports():
    assert portlist_ports(bytes([0x80, 0x01, 0x00, 0x40])) == {1, 16, 26}
    assert portlist_ports(labnet.portlist([1, 8, 9, 289])) == {1, 8, 9, 289}
    assert portlist_ports(b"") == set() and portlist_ports(None) == set()


def test_graph_edges_and_exports(tmp_path):
    inv, _, _ = crawl()
    g = build_graph(inv)
    kinds = {}
    for u, v, a in g.edges(data=True):
        kinds.setdefault(a["kind"], []).append((u, v, a))
    l2 = {tuple(sorted((u, v))) for u, v, _ in kinds["cdp"] + kinds["lldp"]}
    assert ("10.0.0.1", "10.0.0.2") in l2  # R1 - SW1 via CDP (deduped both directions)
    assert ("10.0.0.2", "10.1.0.2") in l2  # SW1 - SW2 via LLDP
    assert ("10.0.0.2", "10.1.0.60") in l2  # SW1 - AP resolved to the ARP host by chassis MAC
    assert ("10.1.0.2", "10.1.0.90") in l2  # SW2 - phone resolved by LLDP mgmt IP
    assert g.nodes["10.1.0.60"]["role"] == "wireless" and g.nodes["10.1.0.60"]["label"] == "ap-lobby"
    assert g.nodes["10.1.0.90"]["role"] == "phone"
    assert "stub:branch-fw" in g and g.nodes["stub:branch-fw"]["role"] == "unpolled"
    # exactly one R1-SW1 edge: L3 edge suppressed when an L2 link exists
    assert sum(1 for _ in g.subgraph(["10.0.0.1", "10.0.0.2"]).edges) == 1
    # FDB host placement: A and B on SW2 ports 3/4, C on SW1 Gi1/0/5, nothing on uplinks
    fdb = {(u, v): a for u, v, a in kinds["fdb"]}
    assert fdb[("10.1.0.2", "10.1.0.50")]["port"] == "3" and fdb[("10.1.0.2", "10.1.0.51")]["port"] == "4"
    assert fdb[("10.0.0.2", "10.2.0.10")]["port"] == "Gi1/0/5" and fdb[("10.0.0.2", "10.2.0.10")]["vlan"] == 20
    assert not any(a["port"] in ("24", "Gi1/0/2") for a in fdb.values())
    # subnet membership
    assert g.has_edge("10.0.0.2", "10.1.0.0/24") and g.has_edge("10.1.0.50", "10.1.0.0/24")
    # exports
    p = tmp_path / "m"
    render_html(g, str(p) + ".html"), export_graphml(g, str(p) + ".graphml"), export_dot(g, str(p) + ".dot")
    files = export_csv(inv, g, str(p) + "-")
    assert (p.parent / "m.html").stat().st_size > 5000 and len(files) == 8
    assert "acc-sw2" in (p.parent / "m-hosts.csv").read_text()  # host A placed on acc-sw2
    assert "USERS" in (p.parent / "m-vlans.csv").read_text()
    import csv

    with open(p.parent / "m-hardware.csv", encoding="utf-8") as f:
        hw = list(csv.DictReader(f))
    assert {"device": "10.0.0.2", "class": "port", "model": "SFP-10G-SR", "serial": "AVD2045K1LM", "fru": "True"}.items() <= next(r for r in hw if r["serial"] == "AVD2045K1LM").items()
    with open(p.parent / "m-interfaces.csv", encoding="utf-8") as f:
        ifs = {(r["device"], r["name"]): r for r in csv.DictReader(f)}
    assert (ifs[("10.0.0.2", "Gi1/0/2")]["mode"], ifs[("10.0.0.2", "Gi1/0/2")]["vlan"]) == ("trunk", "1")
    assert ifs[("10.1.0.2", "21")]["lag"] == "Trk1" and ifs[("10.0.0.2", "Vlan10")]["vlan"] == ""
    with open(p.parent / "m-devices.csv", encoding="utf-8") as f:
        assert {r["ip"]: r["os_version"] for r in csv.DictReader(f)}["10.0.0.2"] == "16.12.4"
    ipam = (p.parent / "m-ipam.csv").read_text()
    assert "10.1.0.0/24" in ipam and "utilisation_pct" in ipam
    txt = text_summary(inv, g)
    assert "core-rtr" in txt and "branch-fw" in txt
    # round-trip persistence
    inv.save(str(p) + ".json")
    inv2 = Inventory.load(str(p) + ".json")
    assert set(inv2.devices) == set(inv.devices) and inv2.devices["10.0.0.2"].vlans == inv.devices["10.0.0.2"].vlans
    assert len(build_graph(inv2).edges) == len(g.edges)


def test_oui_vendor_and_enrichment():
    # Bundled IEEE table: a MAC alone should name the organisation that owns the block.
    assert oui_vendor("00:50:56:11:22:33") == "VMware"
    assert oui_vendor("b8-27-eb-00-11-22").startswith("Raspberry Pi")
    assert oui_vendor("005056112233") == "VMware"
    assert oui_vendor("zz:zz") == "" and oui_vendor(None) == "" and oui_vendor("de:ad:be:ef:00:0a") == ""
    # An ARP-only host has a MAC and nothing else; enrichment must still type it.
    inv = Inventory()
    inv.touch_host("10.9.9.9", "arp", "00:50:56:11:22:33")
    enrich_inventory(inv)
    h = inv.hosts["10.9.9.9"]
    assert h.vendor == "VMware" and h.role == "vm"


def test_ipam_and_vlan_rows():
    inv, _, _ = crawl()
    build_graph(inv)  # fills subnet gateways
    rows = {r["cidr"]: r for r in ipam_rows(inv)}
    users = rows["10.1.0.0/24"]
    assert users["size"] == 256 and users["usable"] == 254
    # SW1 (10.1.0.1, deduped onto 10.0.0.2), SW2 and the ARP/LLDP hosts all count as in use
    assert users["used"] >= 4 and users["free"] == users["usable"] - users["used"]
    assert 0 < users["utilisation_pct"] < 100 and users["swept"] is False
    assert "dist-sw1" in users["gateways"]
    p2p = rows["10.0.0.0/30"]
    assert p2p["usable"] == 2 and p2p["used"] == 2 and p2p["utilisation_pct"] == 100.0
    vl = vlan_rows(inv)
    assert vl[10][0] == {"USERS"} and "dist-sw1" in vl[20][1]


def test_xlsx_report(tmp_path):
    from openpyxl import load_workbook

    inv, _, _ = crawl()
    g = build_graph(inv)
    out = tmp_path / "acme.xlsx"
    export_xlsx(inv, g, str(out))
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Devices", "IPAM", "VLANs", "Links", "Hosts", "Interfaces", "Hardware", "Findings", "Compliance", "Hardware support", "Dependencies", "Gaps"]
    devices = list(wb["Devices"].values)
    assert devices[0][0] == "IP" and any(r[1] == "core-rtr" for r in devices[1:])
    assert devices[0][4:6] == ("Model", "OS version") and any(r[1] == "acc-sw2" and r[5] == "YA.16.10.0016" for r in devices[1:])
    ifs = list(wb["Interfaces"].values)
    col = {h: n for n, h in enumerate(ifs[0])}
    gi5 = next(r for r in ifs[1:] if r[0] == "10.0.0.2" and r[col["Name"]] == "Gi1/0/5")
    assert (gi5[col["VLAN"]], gi5[col["Mode"]], gi5[col["LAG"]]) == (20, "access", None)  # openpyxl reads "" back as None
    hw = list(wb["Hardware"].values)
    assert hw[0] == ("Device", "Device name", "Class", "Name", "Description", "Model", "Serial", "HW rev", "FW rev", "SW rev", "FRU")
    assert ("10.0.0.2", "dist-sw1", "powerSupply", "Switch 2 - Power Supply A") == next(r for r in hw if r[6] == "LIT21330DEF")[:4]
    assert sum(r[0] == "10.0.0.2" for r in hw[1:]) == 8 and sum(r[0] == "10.0.0.1" for r in hw[1:]) == 1
    assert any(r[0] == "10.0.0.0/30" for r in list(wb["IPAM"].values)[1:])
    # the gaps sheet is the point of the pack: the firewall nobody gave us credentials for
    assert any("branch-fw" in str(r[1]) for r in list(wb["Gaps"].values)[1:])
    assert wb["Devices"].freeze_panes == "A2" and wb["Devices"].auto_filter.ref.startswith("A1:")


def test_target_file_and_scope():
    import tempfile

    from netmap.cli import build_parser, read_target_file, scope_from, targets_from

    with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False) as f:
        f.write("# ranges handed over by the target\n10.20.0.0/24\n10.21.0.0/24, 10.22.0.0/24\n\n192.168.5.10\nnot-a-subnet\n")
        path = f.name
    assert read_target_file(path) == ["10.20.0.0/24", "10.21.0.0/24", "10.22.0.0/24", "192.168.5.10/32"]
    args = build_parser().parse_args(["crawl", "--target", "10.30.0.0/24", "--target-file", path])
    targets = targets_from(args, {})
    assert [str(t) for t in targets] == ["10.30.0.0/24", "10.20.0.0/24", "10.21.0.0/24", "10.22.0.0/24", "192.168.5.10/32"]
    # no --scope: the targets are the scope, so nothing else is touched
    scope, exclude = scope_from(args, {}, targets)
    assert [str(s) for s in scope] == [str(t) for t in targets] and exclude == []
    assert in_scope("10.30.0.5", scope, exclude) and not in_scope("10.31.0.5", scope, exclude)
    # an explicit wider scope is kept, and a target outside it is still added rather than dropped
    args2 = build_parser().parse_args(["crawl", "--target", "172.16.4.0/24", "--scope", "10.0.0.0/8", "--exclude", "10.30.0.0/24"])
    t2 = targets_from(args2, {})
    scope2, exclude2 = scope_from(args2, {}, t2)
    assert [str(s) for s in scope2] == ["10.0.0.0/8", "172.16.4.0/24"]
    assert in_scope("172.16.4.9", scope2, exclude2) and not in_scope("10.30.0.9", scope2, exclude2)
    # config file can carry the same lists
    t3 = targets_from(build_parser().parse_args(["crawl"]), {"crawl": {"targets": ["10.40.0.0/24"], "target_files": [path]}})
    assert str(t3[0]) == "10.40.0.0/24" and len(t3) == 5
    os.unlink(path)


def test_probe_all_expands_targets_without_pinging():
    import ipaddress as ia

    from netmap.sweep import discover_targets

    inv = Inventory()
    scope = [ia.ip_network("10.50.0.0/24")]
    ips = asyncio.run(discover_targets(inv, [ia.ip_network("10.50.0.0/30")], scope, [], probe_all=True))
    assert ips == ["10.50.0.1", "10.50.0.2"]  # network and broadcast excluded
    assert "10.50.0.0/30" in inv.subnets and "target" in inv.subnets["10.50.0.0/30"].sources
    # an entered range is probed in full whatever its size
    assert len(asyncio.run(discover_targets(inv, [ia.ip_network("10.50.0.0/24")], scope, [], probe_all=True))) == 254
    # exclusions win inside a target
    ips2 = asyncio.run(discover_targets(inv, [ia.ip_network("10.50.0.0/30")], scope, [ia.ip_network("10.50.0.2/32")], probe_all=True))
    assert ips2 == ["10.50.0.1"]


def test_outputs_are_utf8_regardless_of_locale(tmp_path):
    """Device-supplied text must survive every export.

    On Windows `open(path, "w")` encodes with the legacy code page, which cannot represent
    the embedded map viewer or a sysDescr with an accent in it - the crawl would finish and
    then die while writing the report. Every writer declares UTF-8; this holds it there.
    """
    inv = Inventory()
    d = Device(id="10.0.0.1", name="sw-café-01", sysdescr="Rôle: distribution — 10 Gb/s ‑ tëst", location="Zürich, 3° étage", vendor="Cisco", role="switch")
    d.interfaces.append(Interface(index=1, name="Gi1/0/1", alias="lien ↔ cœur", ips=["10.0.0.1/24"]))
    d.ips.append("10.0.0.1")
    inv.add_device(d)
    inv.touch_host("10.0.0.50", "arp", "00:50:56:11:22:33")
    g = build_graph(inv)
    p = tmp_path / "u"
    render_html(g, f"{p}.html")
    export_dot(g, f"{p}.dot")
    export_graphml(g, f"{p}.graphml")
    export_xlsx(inv, g, f"{p}.xlsx")
    files = export_csv(inv, g, f"{p}-")
    inv.save(f"{p}.json")
    for path in [f"{p}.dot", f"{p}.graphml", *files]:
        open(path, encoding="utf-8").read()  # raises UnicodeDecodeError if mis-encoded
    for path in [f"{p}.dot", f"{p}-devices.csv", f"{p}-interfaces.csv"]:
        assert "sw-café-01" in open(path, encoding="utf-8").read(), path
    # The map escapes the data as \uXXXX, so what made Windows fail is the embedded viewer:
    # the file has to be written as UTF-8 whatever the console code page says.
    html = open(f"{p}.html", encoding="utf-8").read()
    assert "sw-caf\\u00e9-01" in html and any(ord(ch) > 127 for ch in html)
    assert "Zürich" in open(f"{p}-devices.csv", encoding="utf-8").read()
    assert "lien ↔ cœur" in open(f"{p}-interfaces.csv", encoding="utf-8").read()
    assert Inventory.load(f"{p}.json").devices["10.0.0.1"].name == "sw-café-01"

    from openpyxl import load_workbook

    assert any(r[1] == "sw-café-01" for r in list(load_workbook(f"{p}.xlsx")["Devices"].values)[1:])


def test_no_text_write_relies_on_the_platform_encoding():
    """Guard the whole package, not just the paths the test above happens to exercise."""
    import pathlib

    offenders = []
    for src in sorted(pathlib.Path(__file__).resolve().parent.parent.joinpath("netmap").glob("*.py")):
        for n, line in enumerate(src.read_text(encoding="utf-8").splitlines(), 1):
            if re.search(r"\bopen\(", line) and '"rb"' not in line and "encoding=" not in line:
                offenders.append(f"{src.name}:{n}: {line.strip()}")
    assert not offenders, "text I/O without an explicit encoding breaks on a non-UTF-8 console:\n" + "\n".join(offenders)


def test_max_depth_and_resume(tmp_path):
    inv, prober, _ = crawl(max_depth=1)
    assert set(inv.devices) == {"10.0.0.1", "10.0.0.2"}  # SW2 is depth 2
    out = tmp_path / "map.json"
    inv.save(str(out))
    # resume with deeper limit: already-collected devices are not re-probed
    lab = labnet.build("10")
    prober2 = make_prober({ip: d.values() for ip, d in lab.items()})
    cfg = CrawlConfig(seeds=["10.0.0.1"], credentials=[Credential(kind="v2c", community="lab")], scope=[ipaddress.ip_network("10.0.0.0/8")], max_depth=5)
    inv2 = Inventory.load(str(out))
    asyncio.run(Crawler(cfg, inv2, engine=object(), prober=prober2).run())
    assert "10.0.0.1" not in prober2.probed
    # SW2 is reachable on resume because the unmatched-neighbour pass re-enqueues LLDP mgmt IPs
    assert "10.1.0.2" in inv2.devices


def test_collects_fhrp_routing_peers_and_stp():
    inv, _, _ = crawl()
    sw1, sw2, r1 = inv.devices["10.0.0.2"], inv.devices["10.1.0.2"], inv.devices["10.0.0.1"]
    hsrp = [g for g in sw1.redundancy if g["proto"] == "hsrp"]
    assert hsrp and hsrp[0]["vip"] == "10.1.0.254" and hsrp[0]["state"] == "active" and hsrp[0]["interface"] == "Vlan10"
    assert {(p["proto"], p["addr"], p["state"]) for p in r1.peers} == {("ospf", "10.0.0.2", "full"), ("bgp", "203.0.113.1", "established")}
    assert any(p["extra"] == "AS65001" for p in r1.peers)
    assert sw1.stp["is_root"] is True and sw1.stp["priority"] == 24576
    assert sw2.stp["is_root"] is False and sw2.stp["root"] == sw1.stp["root"] and sw2.stp["root_port"] == "24"


def test_interface_counters_poe_and_utilisation():
    from netmap.collect import apply_counter_deltas
    from netmap.model import Device, Interface

    inv, _, _ = crawl()
    sw1 = inv.devices["10.0.0.2"]
    gi2 = sw1.iface(2)
    assert gi2.in_errors == 1200 and gi2.duplex == "half"
    p24 = sw1.iface(24)
    assert p24.poe_status == "delivering" and p24.poe_watts == 25.5 and p24.poe_class == "3"
    assert sw1.iface(5).poe_status == "searching"
    assert sw1.poe_budget_w == 740 and sw1.poe_used_w == 130

    # utilisation and error rate come from two counter snapshots 10s apart
    old = Device(id="1")
    old.interfaces = [Interface(index=1, speed_mbps=1000, in_octets=0, out_octets=0, in_errors=0, out_errors=0, counters_at=1000.0)]
    new = Device(id="1")
    new.interfaces = [Interface(index=1, speed_mbps=1000, in_octets=1_250_000_000, out_octets=0, in_errors=50, out_errors=0, counters_at=1010.0)]
    apply_counter_deltas(old, new)
    # 1.25 GB in 10 s = 1 Gbit/s = 100% of a 1 Gbps link
    assert new.interfaces[0].in_util_pct == 100.0 and new.interfaces[0].err_rate == 5.0
    # counter wrap (delta negative) is ignored, not shown as negative util
    new.interfaces[0].in_octets = 0
    apply_counter_deltas(old, new)
    assert new.interfaces[0].in_util_pct >= 0
