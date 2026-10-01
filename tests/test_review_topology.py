"""Review fixes in the topology/model/views layer: graph cache, gateways, host access over
LLDP, findings, query parsing, diff/reconcile matching, export hygiene and the layered
layout under load. The sample campus is loaded read-only; synthetic inventories cover the
cases it does not have."""
import csv
import re

import pytest

from subnetsleuth import diagram, layout
from subnetsleuth.compliance import _default_stp_priority, compliance_checks
from subnetsleuth.diff import compare
from subnetsleuth.graph import _is_uplink_iface, build_graph, export_csv, stub_endpoint_role
from subnetsleuth.model import Device, FdbEntry, Host, Interface, Inventory, Neighbor, Route
from subnetsleuth.paths import host_access, path_to, route_path
from subnetsleuth.query import QueryError, run_query
from subnetsleuth.util import resource_path
from subnetsleuth.views import DEVICE_COLUMNS, Snapshot, device_rows, finding_rows, vlan_rows_view

SAMPLE = resource_path("data", "sample-campus.sleuth")


@pytest.fixture(scope="module")
def campus():
    return Inventory.load(SAMPLE)


# ---------------------------------------------------------------- helpers
def _switch(ip, name, role="switch", **kw):
    d = Device(id=ip, name=name, role=role, ips=[ip], **kw)
    d.interfaces.append(Interface(index=99, name="Vlan99", ips=[f"{ip}/24"], type=53, admin_up=True, oper_up=True))
    return d


def _link(a, ai, ap, b, bi, bp, proto="lldp"):
    a.neighbors.append(Neighbor(proto=proto, local_if_index=ai, local_port=ap, remote_name=b.name, remote_port=bp, remote_mgmt_ips=[b.id]))
    b.neighbors.append(Neighbor(proto=proto, local_if_index=bi, local_port=bp, remote_name=a.name, remote_port=ap, remote_mgmt_ips=[a.id]))


def _small_net():
    """core (L3) - two access switches, a few hosts on ports, one subnet each."""
    inv = Inventory()
    core = Device(id="10.0.0.1", name="core", role="l3switch", ips=["10.0.0.1"])
    core.interfaces = [Interface(index=1, name="Vlan10", ips=["10.10.0.1/24"], type=53, admin_up=True, oper_up=True),
                       Interface(index=2, name="Vlan99", ips=["10.0.0.1/24"], type=53, admin_up=True, oper_up=True),
                       Interface(index=11, name="Te1/0/1", mode="trunk", type=6, admin_up=True, oper_up=True),
                       Interface(index=12, name="Te1/0/2", mode="trunk", type=6, admin_up=True, oper_up=True)]
    core.routes = [Route(dest="0.0.0.0/0", nexthop="10.0.0.254"), Route(dest="10.10.0.0/24", nexthop="0.0.0.0", type=3)]
    core.stp = {"root": "aa:aa", "priority": 4096, "is_root": True}
    sw1 = _switch("10.0.0.11", "sw1")
    sw2 = _switch("10.0.0.12", "sw2")
    for s in (sw1, sw2):
        s.interfaces += [Interface(index=i, name=f"Gi1/0/{i}", mode="access", vlan=10, type=6, admin_up=True, oper_up=(i < 4)) for i in range(1, 9)]
        s.interfaces.append(Interface(index=49, name="Te1/1/1", mode="trunk", vlan=1, type=6, admin_up=True, oper_up=True))
        s.routes = [Route(dest="0.0.0.0/0", nexthop="10.0.0.1")]
        s.stp = {"root": "aa:aa", "priority": 32769, "is_root": False}
    _link(core, 11, "Te1/0/1", sw1, 49, "Te1/1/1")
    _link(core, 12, "Te1/0/2", sw2, 49, "Te1/1/1")
    for d in (core, sw1, sw2):
        inv.add_device(d)
    n = 0
    for sw in (sw1, sw2):
        for port in range(1, 4):
            n += 1
            mac = f"02:00:00:00:00:{n:02x}"
            ip = f"10.10.0.{n + 10}"
            h = inv.touch_host(ip, "arp", mac)
            h.hostname = f"pc{n}"
            sw.fdb.append(FdbEntry(mac=mac, if_index=port, vlan=10))
            core.fdb.append(FdbEntry(mac=mac, if_index=11 if sw is sw1 else 12, vlan=10))
    return inv


# ---------------------------------------------------------------- 1/2: cache and revisions
def test_graph_cache_hits_and_invalidates(campus):
    g = build_graph(campus)
    assert build_graph(campus) is g
    assert build_graph(campus, include_hosts=False) is not g  # a different view is a different key
    campus.annotate("10.99.0.2", notes="checked")
    g2 = build_graph(campus)
    assert g2 is not g and build_graph(campus) is g2
    campus.annotate("10.99.0.2", notes="")


def test_rev_bumps_on_every_graph_changing_mutation():
    inv = Inventory()
    r = inv.rev
    inv.touch_host("10.1.1.5", "arp", "02:00:00:00:00:05")
    assert inv.rev > r
    r = inv.rev
    inv.add_subnet("10.1.1.0/24", "device")
    assert inv.rev > r
    r = inv.rev
    inv.add_subnet("10.1.1.0/24", "device")  # nothing new
    assert inv.rev == r
    inv.add_device(Device(id="10.1.1.1", name="r1"))
    assert inv.rev > r
    r = inv.rev
    assert inv.remove_host("10.1.1.5") is True and inv.rev > r and "10.1.1.5" not in inv.hosts
    assert inv.remove_host("10.1.1.5") is False
    r = inv.rev
    assert inv.remove_device("10.1.1.1") is True and inv.rev > r and "10.1.1.1" not in inv.ip_to_device


def test_remove_host_invalidates_graph_and_lpm_cache_survives_remove_add():
    inv = _small_net()
    g = build_graph(inv)
    assert "10.10.0.11" in g
    inv.remove_host("10.10.0.11")
    g2 = build_graph(inv)
    assert g2 is not g and "10.10.0.11" not in g2
    # longest-prefix cache: remove one subnet, add another of the same count -> not stale
    assert inv.subnet_for_ip("10.10.0.50") == "10.10.0.0/24"
    inv.remove_subnet("10.10.0.0/24")
    inv.add_subnet("10.20.0.0/24", "target")
    assert inv.subnet_for_ip("10.10.0.50") is None and inv.subnet_for_ip("10.20.0.7") == "10.20.0.0/24"


# ---------------------------------------------------------------- 4: LAG names
@pytest.mark.parametrize("name,lag", [("Port 1", False), ("port 24", False), ("Port24", False), ("Po1", True), ("po10", True),
                                      ("Port-channel1", True), ("Port-Channel 2", True), ("ae0", True), ("bond0", True), ("lag3", True), ("Team1", True),
                                      ("GigabitEthernet1/0/1", False), ("Portland-uplink", False)])
def test_lag_name_pattern(name, lag):
    d = Device(id="1.1.1.1", interfaces=[Interface(index=1, name=name, type=6)])
    assert _is_uplink_iface(d, 1) is lag


# ---------------------------------------------------------------- 5: gateways
def test_gateways_are_routers_not_every_member(campus):
    build_graph(campus)
    mgmt = campus.subnets["10.99.0.0/24"]
    assert set(mgmt.gateways) == {"10.99.0.2", "10.99.0.3"}  # the two cores, not seven access switches
    assert "10.99.0.11" not in mgmt.gateways
    # rebuilt each time, never appended forever
    build_graph(campus, include_hosts=False)
    assert len(campus.subnets["10.99.0.0/24"].gateways) == 2


def test_gateway_by_routes_and_fhrp():
    inv = _small_net()
    sw1 = inv.devices["10.0.0.11"]
    sw1.routes.append(Route(dest="192.168.5.0/24", nexthop="10.0.0.99"))  # a "switch" that routes for others
    sw2 = inv.devices["10.0.0.12"]
    sw2.redundancy.append({"proto": "vrrp", "group": "1", "vip": "10.0.0.250", "state": "master", "priority": 110})
    build_graph(inv)
    assert set(inv.subnets["10.0.0.0/24"].gateways) == {"10.0.0.1", "10.0.0.11", "10.0.0.12"}


# ---------------------------------------------------------------- 3: host access via LLDP
def test_path_to_lldp_announced_host_ends_at_access_port(campus):
    g = build_graph(campus)
    assert host_access(campus, g, "10.99.0.110") == ("10.99.0.11", "Gi1/0/1")
    p = path_to(campus, g, "10.99.0.110")
    assert p.ok and [h.node for h in p.hops] == ["10.99.0.2", "10.99.0.11", "10.99.0.110"]
    assert p.hops[-1].in_port == "Gi1/0/1" and "Gi1/0/1" in p.hops[-1].detail


def test_path_to_fdb_host_still_works():
    inv = _small_net()
    g = build_graph(inv)
    p = path_to(inv, g, "10.10.0.14", origin="10.0.0.1")
    assert p.ok and p.hops[-2].node == "10.0.0.12" and p.hops[-1].in_port == "Gi1/0/1"


# ---------------------------------------------------------------- 14: route_path
def test_route_path_ignores_default_address_and_reports_loops():
    inv = Inventory()
    a = Device(id="10.0.0.1", name="a", role="router", ips=["10.0.0.1"])
    a.interfaces = [Interface(index=1, name="Tunnel0", ips=["0.0.0.0/0"]), Interface(index=2, name="Lo0", ips=["127.0.0.1/8"]),
                    Interface(index=3, name="Gi0/0", ips=["10.0.0.1/30"])]
    a.routes = [Route(dest="10.50.0.0/16", nexthop="10.0.0.2", if_index=3)]
    b = Device(id="10.0.0.2", name="b", role="router", ips=["10.0.0.2"])
    b.interfaces = [Interface(index=1, name="Gi0/0", ips=["10.0.0.2/30"])]
    b.routes = [Route(dest="10.50.0.0/16", nexthop="10.0.0.1", if_index=1)]  # points back: a loop
    inv.add_device(a)
    inv.add_device(b)
    hops = route_path(inv, "10.0.0.1", "10.50.1.1")
    assert not any("directly connected" in h.detail for h in hops)
    assert "routing loop" in hops[-1].detail and hops[-1].node == "10.0.0.1"
    # hop limit is reported rather than silently truncated
    hops = route_path(inv, "10.0.0.1", "10.50.1.1", max_hops=1)
    assert "hop limit" in hops[-1].detail


# ---------------------------------------------------------------- 7: VLAN MAC counts
def test_vlan_hosts_learned_counts_distinct_macs():
    inv = _small_net()  # every MAC is in the FDB of its access switch *and* the core
    rows = {r["vlan"]: r for r in vlan_rows_view(Snapshot(inv))}
    inv.devices["10.0.0.1"].vlans[10] = "users"
    rows = {r["vlan"]: r for r in vlan_rows_view(Snapshot(inv))}
    assert rows[10]["hosts"] == 6  # not 12


# ---------------------------------------------------------------- 10: compliance
def test_stp_default_priority_check():
    assert _default_stp_priority(32768) and _default_stp_priority(32769) and _default_stp_priority(32778)
    assert not _default_stp_priority(0) and not _default_stp_priority(4096) and not _default_stp_priority(None) and not _default_stp_priority(28672)


def test_compliance_http_and_snmp_version():
    inv = _small_net()
    core = inv.devices["10.0.0.1"]
    core.mgmt = {"telnet": True, "http": True}
    core.stp = {"root": "x", "priority": 32769, "is_root": True}
    core.snmp_version = "v2c"
    core.credential = "site-ro"  # label says nothing about the version
    sw1 = inv.devices["10.0.0.11"]
    sw1.credential = "hq-snmpv3"
    sw2 = inv.devices["10.0.0.12"]
    sw2.credential = "branch v2c"
    cats = {}
    for c in compliance_checks(Snapshot(inv)):
        cats.setdefault(c.category, set()).add(c.node)
    assert "10.0.0.1" in cats["Telnet enabled"] and "10.0.0.1" in cats["Cleartext web management"]
    assert "10.0.0.1" in cats["Spanning-tree root by default"]
    assert cats["SNMPv2c in use"] == {"10.0.0.1", "10.0.0.12"}


# ---------------------------------------------------------------- 11: query
def test_query_quoted_values_with_operators_and_joiners(campus):
    s = Snapshot(campus)
    _, rows = run_query(s, 'devices where name ~ "core" and site ~ "comms room"')
    assert rows and all("core" in r["name"] for r in rows)
    _, rows = run_query(s, 'hosts where name ~ "a=b and c or d"')
    assert rows == []
    _, rows = run_query(s, "devices where name = 'core-sw-01' or name = 'core-sw-02'")
    assert {r["name"] for r in rows} == {"core-sw-01", "core-sw-02"}


def test_query_port_matches_open_ports_not_switch_port(campus):
    s = Snapshot(campus)
    _, rows = run_query(s, "hosts where port = 3389")
    assert rows and all(re.search(r"\b3389\b", r["services"]) for r in rows)
    _, rows = run_query(s, "hosts where port = 1")  # Gi1/0/1 must not count as port 1
    assert all(re.search(r"(^|\s)1/", r["services"]) for r in rows)
    _, rows = run_query(s, "dependencies where port > 1000")  # a real numeric column elsewhere
    assert all(int(r["port"]) > 1000 for r in rows)


def test_query_ordering_by_column_kind(campus):
    s = Snapshot(campus)
    _, rows = run_query(s, "hosts where ip > 10.50.0.0 and ip < 10.50.255.255")
    assert rows and all(r["ip"].startswith("10.50.") for r in rows)
    _, rows = run_query(s, "hosts where ip > 10.9.0.0")  # as text "10.9" > "10.50"; as an address it is not
    assert any(r["ip"].startswith("10.50.") for r in rows)
    _, rows = run_query(s, "devices where uptime > 30d")
    assert rows and all(int(r["uptime"]) > 30 * 86400 for r in rows)
    _, rows = run_query(s, "subnets where cidr < 10.20.0.0/16")
    assert rows and all(r["cidr"].startswith(("10.0.", "10.10.")) for r in rows)


def test_query_bad_start_is_a_query_error(campus):
    s = Snapshot(campus)
    with pytest.raises(QueryError):
        run_query(s, "= devices")
    with pytest.raises(QueryError):
        run_query(s, "hosts where port > 5")


# ---------------------------------------------------------------- 12: diff
def test_diff_ignores_placeholder_serials():
    old, new = Inventory(), Inventory()
    old.add_device(Device(id="10.0.0.1", name="a", serial="N/A"))
    old.add_device(Device(id="10.0.0.2", name="b", serial="SHARED1"))
    old.add_device(Device(id="10.0.0.3", name="c", serial="SHARED1"))
    new.add_device(Device(id="10.0.0.9", name="a2", serial="N/A"))
    new.add_device(Device(id="10.0.0.8", name="b2", serial="SHARED1"))
    new.add_device(Device(id="10.0.0.7", name="real", serial="FDO12345"))
    old.add_device(Device(id="10.0.0.4", name="real", serial="fdo12345 "))
    d = compare(old, new)
    moved = d.of("device", "moved")
    assert [c.item for c in moved] == ["10.0.0.7"]  # only the genuine serial
    assert {c.item for c in d.of("device", "added")} == {"10.0.0.9", "10.0.0.8"}


def test_diff_readdressed_hosts_and_uptime_wrap():
    from subnetsleuth.diff import UPTIME_WRAP_S

    old, new = Inventory(), Inventory()
    a = Device(id="10.0.0.1", name="a", uptime_s=496 * 86400, collected_at=1_000_000)
    wrapped = int(496 * 86400 + 5 * 86400 - UPTIME_WRAP_S) + 30  # the counter went round, the device did not restart
    b = Device(id="10.0.0.1", name="a", uptime_s=wrapped, collected_at=1_000_000 + 5 * 86400)
    old.add_device(a)
    new.add_device(b)
    old.hosts["10.1.0.5"] = Host(ip="10.1.0.5", mac="02:00:00:00:00:01", hostname="laptop")
    old.hosts["10.1.0.6"] = Host(ip="10.1.0.6", mac="02:00:00:00:00:02")
    new.hosts["10.1.0.77"] = Host(ip="10.1.0.77", mac="02:00:00:00:00:01")
    new.hosts["10.1.0.6"] = Host(ip="10.1.0.6", mac="02:00:00:00:00:02")
    d = compare(old, new)
    assert not any("rebooted" in c.detail for c in d.of("device"))
    re_ = d.of("host", "readdressed")
    assert len(re_) == 1 and re_[0].item == "10.1.0.77" and "10.1.0.5" in re_[0].detail and re_[0].name == "laptop"
    assert not d.of("host", "added") and not d.of("host", "removed")
    # a real reboot: uptime shorter than the time between the scans
    new.devices["10.0.0.1"].uptime_s = 3600
    assert any("rebooted" in c.detail for c in compare(old, new).of("device"))


# ---------------------------------------------------------------- 13: reconcile
def test_reconcile_normalises_serial_and_accepts_annotated_name():
    from subnetsleuth.reconcile import reconcile

    inv = _small_net()
    inv.devices["10.0.0.1"].serial = "FDO2233ABCD"
    inv.annotate("10.0.0.1", name="CORE-A")
    h = inv.touch_host("10.10.0.99", "arp", "02:00:00:00:00:99")
    h.snmp_failed = True
    cols = {"ip": 0, "name": 1, "serial": 2}
    rec = reconcile(inv, [["10.0.0.1", "CORE-A", " fdo2233abcd "], ["10.10.0.99", "old-switch", "XYZ123"], ["10.0.0.11", "sw1", ""]], cols)
    by_ip = {m.listed["ip"]: m for m in rec.matches}
    assert by_ip["10.0.0.1"].differences == []
    assert any("no SNMP answer" in x for x in by_ip["10.10.0.99"].differences)
    assert by_ip["10.0.0.11"].differences == []


# ---------------------------------------------------------------- 16: export hygiene
def test_spreadsheet_cells_starting_with_formula_characters_are_text(tmp_path):
    from openpyxl import load_workbook

    from subnetsleuth.report import export_xlsx

    inv = _small_net()
    inv.devices["10.0.0.1"].location = "=HYPERLINK(\"http://x\")"
    inv.devices["10.0.0.11"].interfaces[0].alias = "+uplink"
    inv.annotate("10.0.0.12", notes="-see ticket 42")
    g = build_graph(inv)
    wb = load_workbook(export_xlsx(inv, g, str(tmp_path / "o.xlsx")))
    formula_cells = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith(("=", "+", "-", "@"))]
    assert formula_cells and all(c.data_type == "s" for c in formula_cells)
    assert {"Findings", "Compliance", "Hardware support", "Dependencies"} <= set(wb.sheetnames)
    hdr = [c.value for c in wb["Devices"][1]]
    assert {"Owner", "Status", "Asset tag", "Notes", "Free ports"} <= set(hdr)
    files = export_csv(inv, g, str(tmp_path / "x_"))
    for p in files:
        with open(p, newline="", encoding="utf-8") as f:
            for row in csv.reader(f):
                assert not any(cell.startswith(("=", "+", "-", "@", "\t")) for cell in row), (p, row)
    with open(tmp_path / "x_devices.csv", encoding="utf-8") as f:
        assert "'=HYPERLINK" in f.read()


def test_render_substitutes_placeholders_once(tmp_path):
    from subnetsleuth.render import render_html

    inv = _small_net()
    inv.annotate("10.0.0.1", notes="__COLORS__ and __DATA__")
    g = build_graph(inv)
    out = tmp_path / "map.html"
    render_html(g, str(out), embed_js=False)
    html = out.read_text(encoding="utf-8")
    assert "__COLORS__ and __DATA__" in html  # the note survived as text
    assert html.count("const COLORS = {") == 1


def test_drawio_double_escapes_html_values_and_tolerates_bad_positions(tmp_path):
    inv = _small_net()
    core = inv.devices["10.0.0.1"]
    core.neighbors[0].local_port = "Te<1>/0/1 & co"
    core.model = "A<b>"
    g = build_graph(inv)
    inv.layout["physical"] = {"10.0.0.1": [1.0, 2.0], "10.0.0.11": ["nan", None], "10.0.0.12": "junk"}
    p = diagram.export_drawio(inv, g, str(tmp_path / "d.drawio"))
    text = open(p, encoding="utf-8").read()
    assert "&amp;lt;1&amp;gt;" in text and "A&amp;lt;b&amp;gt;" in text
    flags = {k: v for k, v in diagram.PRESETS["physical"].items() if k != "title"}
    nodes, edges = diagram.select(g, flags, "physical")
    pos = diagram.positions(nodes, edges, inv.layout["physical"])
    assert pos["10.0.0.1"] == (1.0, 2.0) and all(n in pos for n in nodes)


# ---------------------------------------------------------------- 6/8/9/17/18/22: findings
def test_endpoint_stub_role_and_finding():
    assert stub_endpoint_role(Neighbor(proto="lldp", remote_caps="telephone,bridge")) == "phone"
    assert stub_endpoint_role(Neighbor(proto="lldp", remote_caps="wlan-ap,router")) == "wireless"
    assert stub_endpoint_role(Neighbor(proto="lldp", remote_caps="station")) == "host"
    assert stub_endpoint_role(Neighbor(proto="cdp", remote_caps="router,switch,igmp")) == ""
    inv = _small_net()
    sw1 = inv.devices["10.0.0.11"]
    sw1.neighbors.append(Neighbor(proto="lldp", local_if_index=5, local_port="Gi1/0/5", remote_name="SEP0011", remote_port="1", remote_caps="telephone,bridge"))
    sw1.neighbors.append(Neighbor(proto="lldp", local_if_index=6, local_port="Gi1/0/6", remote_name="dist-old", remote_port="1", remote_caps="bridge,router"))
    rows = finding_rows(Snapshot(inv))
    by_item = {(r["category"], r["item"]): r for r in rows}
    assert by_item[("Endpoint announced over LLDP, address not seen", "SEP0011")]["severity"] == "Info"
    assert by_item[("Neighbour not polled", "dist-old")]["severity"] == "Attention"


def test_subnet_gap_ipv6_and_shared_mac_findings():
    inv = _small_net()
    core = inv.devices["10.0.0.1"]
    core.routes += [Route(dest="10.200.0.0/31", nexthop="10.0.0.2"), Route(dest="10.201.0.0/24", nexthop="10.0.0.2"),
                    Route(dest="10.0.0.0/8", nexthop="10.0.0.2"), Route(dest="10.202.0.0/24", nexthop="10.0.0.2")]
    core.interfaces.append(Interface(index=7, name="Gi1/0/7", ips=["10.202.0.1/32"]))  # a device address inside 10.202.0.0/24
    inv.hosts["fe80::1"] = Host(ip="fe80::1", sources=["mdns"])
    # a next-hop MAC across three hosts in two subnets is dropped and reported; aliases in one subnet are kept
    inv.add_subnet("10.30.0.0/24", "target")
    for ip in ("10.10.0.201", "10.10.0.202", "10.30.0.5"):
        inv.touch_host(ip, "arp", "02:99:99:99:99:99")
    for ip in ("10.10.0.211", "10.10.0.212", "10.10.0.213"):
        inv.touch_host(ip, "arp", "02:11:11:11:11:11")
    rows = finding_rows(Snapshot(inv))
    gaps = {r["item"] for r in rows if r["category"] == "Subnet not yet scanned"}
    assert gaps == {"10.201.0.0/24"}
    assert not any(r["category"] == "Address outside every known subnet" and r["item"] == "fe80::1" for r in rows)
    shared = [r for r in rows if r["category"] == "MAC seen on several addresses"]
    assert len(shared) == 1 and shared[0]["item"] == "02:99:99:99:99:99"
    assert inv.hosts["10.10.0.211"].mac == "02:11:11:11:11:11" and inv.hosts["10.30.0.5"].mac is None


def test_new_structural_findings_and_free_ports():
    inv = _small_net()
    core = inv.devices["10.0.0.1"]
    sw1, sw2 = inv.devices["10.0.0.11"], inv.devices["10.0.0.12"]
    # overlapping ranges from different devices
    sw2.interfaces.append(Interface(index=77, name="Vlan77", ips=["10.10.0.129/25"], type=53))
    inv.add_subnet("10.10.0.128/25", "device")  # what add_device/replace_device do for a collected interface
    # native VLAN mismatch across the core-sw1 cable
    core.iface(11).vlan = 1
    sw1.iface(49).vlan = 99
    # STP: sw2 thinks it is the root at default priority, and names a different root bridge
    sw2.stp = {"root": "bb:bb", "priority": 32769, "is_root": True}
    inv.reindex()
    s = Snapshot(inv)
    rows = finding_rows(s)
    cats = {}
    for r in rows:
        cats.setdefault(r["category"], []).append(r)
    assert {r["item"] for r in cats["Single uplink"]} == {"sw1", "sw2"}
    assert cats["Overlapping subnets"][0]["item"] == "10.10.0.128/25" and "10.10.0.0/24" in cats["Overlapping subnets"][0]["detail"]
    assert cats["Spanning-tree root is an access switch"][0]["item"] == "sw2"
    assert "bb:bb" in cats["Switches disagree on the spanning-tree root"][0]["detail"]
    vm = cats["VLAN mismatch on link"]
    assert len(vm) == 1 and "1 vs 99" in vm[0]["detail"] or "99 vs 1" in vm[0]["detail"]
    # free ports: 8 access ports admin-up, 3 with link -> 5 free; the column is declared for the table
    assert "free_ports" in {c.key for c in DEVICE_COLUMNS}
    assert {r["name"]: r["free_ports"] for r in device_rows(s)}["sw1"] == 5


def test_pvid_only_from_access_ports_and_fdb_vlan_index():
    from subnetsleuth.graph import subnet_vlans

    inv = _small_net()
    core = inv.devices["10.0.0.1"]
    core.interfaces.append(Interface(index=30, name="Gi1/0/30", ips=["10.77.0.1/30"], vlan=1, mode=""))  # routed port, PVID 1
    core.interfaces.append(Interface(index=31, name="Gi1/0/31", ips=["10.78.0.1/24"], vlan=78, mode="access"))
    inv.reindex()
    v = subnet_vlans(inv)
    assert "10.77.0.0/30" not in v or 1 not in v["10.77.0.0/30"]
    assert v["10.78.0.0/24"] == {78}
    g = build_graph(inv)
    assert all(a.get("vlan") == 10 for _, _, a in g.edges(data=True) if a.get("kind") == "fdb")


def test_neighbour_short_names_only_match_when_unique():
    inv = Inventory()
    a = Device(id="10.0.0.1", name="sw1.site-a.example", role="switch", ips=["10.0.0.1"])
    b = Device(id="10.0.0.2", name="sw1.site-b.example", role="switch", ips=["10.0.0.2"])
    c = Device(id="10.0.0.3", name="core", role="l3switch", ips=["10.0.0.3"])
    c.neighbors.append(Neighbor(proto="lldp", local_if_index=1, local_port="1", remote_name="sw1.site-b.example", remote_port="49"))
    c.neighbors.append(Neighbor(proto="lldp", local_if_index=2, local_port="2", remote_name="sw1", remote_port="49"))
    for d in (a, b, c):
        inv.add_device(d)
    assert inv.device_for_name("sw1.site-b.example").id == "10.0.0.2"
    assert inv.device_for_name("sw1") is None  # ambiguous
    assert inv.device_for_name("CORE").id == "10.0.0.3"
    g = build_graph(inv, include_hosts=False, include_subnets=False)
    assert g.has_edge("10.0.0.3", "10.0.0.2") and not g.has_edge("10.0.0.3", "10.0.0.1")
    assert "stub:sw1" in g


# ---------------------------------------------------------------- 20: layout under load
def test_layered_layout_packs_hosts_under_switches_with_subnets_shown():
    import time

    inv = Inventory()
    core = Device(id="10.0.0.1", name="core", role="l3switch", ips=["10.0.0.1"])
    core.interfaces = [Interface(index=1, name="Vlan10", ips=["10.10.0.1/16"], type=53), Interface(index=2, name="Vlan99", ips=["10.0.0.1/24"], type=53)]
    inv.add_device(core)
    n = 0
    for s in range(12):
        sw = _switch(f"10.0.0.{s + 10}", f"acc-{s:02d}")
        sw.interfaces.append(Interface(index=49, name="Te1/1/1", mode="trunk", type=6))
        _link(core, 100 + s, f"Te1/0/{s + 1}", sw, 49, "Te1/1/1")
        for k in range(120):
            n += 1
            mac = f"02:00:00:{(n >> 8) & 255:02x}:{n & 255:02x}:01"
            ip = f"10.10.{n // 250}.{n % 250 + 2}"
            inv.touch_host(ip, "arp", mac)
            sw.fdb.append(FdbEntry(mac=mac, if_index=1 + k % 48, vlan=10))
        inv.add_device(sw)
    g = build_graph(inv)
    flags = {**{k: v for k, v in diagram.PRESETS["all"].items() if k != "title"}, "hosts": True}
    nodes, edges = diagram.select(g, flags, "all")
    assert sum(1 for a in nodes.values() if a.get("kind") == "host") == 12 * 120
    t0 = time.time()
    pos = layout.layered(nodes, edges)
    assert time.time() - t0 < 3.0
    xs = [x for x, _ in pos.values()]
    ys = [y for _, y in pos.values()]
    width, height = max(xs) - min(xs), max(ys) - min(ys)
    assert width < 12 * 120 * 104 / 4  # far narrower than one row of 1440 hosts
    assert height > 400  # blocks wrapped into several rows
    # each host sits in the block under its switch: same x band, below it
    sw_x = {n: pos[n][0] for n, a in nodes.items() if a.get("kind") == "device" and a.get("role") == "switch"}
    for u, v, a in edges:
        if a.get("kind") == "fdb":
            sw, host = (u, v) if u in sw_x else (v, u)
            assert pos[host][1] > pos[sw][1]
    # plain (u, v) pairs still work
    assert set(layout.layered(nodes, [(u, v) for u, v, _ in edges])) == set(nodes)
